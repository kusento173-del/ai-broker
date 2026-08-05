from __future__ import annotations

import json
import os
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import requests

from .config import dimensions
from .paths import WorkDirs, today_key
from .utils import display_value, row_value, safe_print, text_value


TABLE_COLUMNS = [
    "分析日期", "分析时间", "后台", "主播昵称", "抖音号/短ID", "对应运营", "直播标题", "开播时长",
    "当前人气", "累计观众", "音浪/流水", "预警原因", "维度", "优先级", "维度结论",
    "证据摘要", "问题诊断", "转化影响", "优化建议", "10分钟内立刻执行", "下次开播前准备", "摘要消息",
]
WIDE_COLUMNS = {"证据摘要", "问题诊断", "转化影响", "优化建议", "10分钟内立刻执行", "下次开播前准备", "摘要消息"}


def local_table_path(config: dict[str, Any], dirs: WorkDirs) -> tuple[Path, str]:
    table = config.get("local_table", {})
    directory = str(table.get("directory", "analysis_table"))
    filename = str(table.get("filename", "预警主播分析_{date}.xlsx")).format(date=today_key())
    sheet = str(table.get("sheet_name", "AI分析"))[:31]
    path = dirs.day() / directory / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    return path, sheet


def table_row(dimension: str, row: dict[str, str], item: dict[str, Any]) -> dict[str, Any]:
    now = datetime.now()
    return {
        "分析日期": now.strftime("%Y-%m-%d"),
        "分析时间": now.strftime("%Y-%m-%d %H:%M:%S"),
        "后台": display_value(row, "后台"),
        "主播昵称": display_value(row, "主播昵称"),
        "抖音号/短ID": display_value(row, "抖音号/短ID"),
        "对应运营": display_value(row, "经纪人"),
        "直播标题": display_value(row, "直播标题"),
        "开播时长": display_value(row, "开播时长"),
        "当前人气": display_value(row, "当前人气"),
        "累计观众": display_value(row, "累计观众"),
        "音浪/流水": display_value(row, "音浪/流水"),
        "预警原因": row_value(row, "预警原因"),
        "维度": dimension,
        "优先级": text_value(item.get("优先级")),
        "维度结论": text_value(item.get("维度结论")),
        "证据摘要": text_value(item.get("证据摘要")),
        "问题诊断": text_value(item.get("问题诊断")),
        "转化影响": text_value(item.get("转化影响")),
        "优化建议": text_value(item.get("优化建议")),
        "10分钟内立刻执行": text_value(item.get("10分钟内立刻执行")),
        "下次开播前准备": text_value(item.get("下次开播前准备")),
        "摘要消息": text_value(item.get("摘要消息")),
    }


def analysis_rows(row: dict[str, str], analysis: dict[str, Any], config: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for dimension in dimensions(config):
        item = analysis.get(dimension, {})
        if not isinstance(item, dict):
            item = {"摘要消息": json.dumps(item, ensure_ascii=False)}
        rows.append(table_row(dimension, row, item))
    return rows


def append_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.create_sheet(sheet_name)
        if sheet.max_row == 0:
            sheet.append(TABLE_COLUMNS)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(TABLE_COLUMNS)

    if sheet.max_row == 1:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    headers = [cell.value for cell in sheet[1] if cell.value]
    for column in TABLE_COLUMNS:
        if column not in headers:
            sheet.cell(row=1, column=len(headers) + 1, value=column)
            headers.append(column)
    for data in rows:
        sheet.append([data.get(column, "") for column in headers])
    for cells in sheet.columns:
        header = cells[0].value
        sheet.column_dimensions[cells[0].column_letter].width = 36 if header in WIDE_COLUMNS else 18
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    temp = path.with_name(f"{path.stem}.tmp{path.suffix}")
    workbook.save(temp)
    temp.replace(path)


def write_local_table(row: dict[str, str], analysis: dict[str, Any], config: dict[str, Any], dirs: WorkDirs) -> Path:
    path, sheet = local_table_path(config, dirs)
    append_xlsx(path, sheet, analysis_rows(row, analysis, config))
    return path


def date_millis(date_text: str) -> int:
    return int(datetime.strptime(date_text, "%Y-%m-%d").timestamp() * 1000)


def table_webhook(config: dict[str, Any]) -> dict[str, Any]:
    data = config.get("table_webhook", {})
    env = str(data.get("url_env", "FEISHU_TABLE_WEBHOOK_URL")).strip()
    url = str(data.get("url", "")).strip() or os.environ.get(env, "").strip()
    if not url:
        raise RuntimeError(f"飞书表格 webhook 未配置，请设置 {env} 或 table_webhook.url。")
    return {"url": url, "interval": float(data.get("submit_interval_seconds", 0.5)), "timeout": float(data.get("timeout_seconds", 30))}


def post_json(url: str, payload: dict[str, Any], timeout: float, error_prefix: str) -> None:
    delays = [3, 10, 30]
    for attempt in range(len(delays) + 1):
        try:
            response = requests.post(url, json=payload, timeout=timeout)
            if response.ok:
                return
            message = f"HTTP {response.status_code}: {response.text[:500]}"
            retry = response.status_code in {429, 500, 502, 503, 504}
        except requests.RequestException as exc:
            message = str(exc)
            retry = True
        if attempt >= len(delays) or not retry:
            raise RuntimeError(f"{error_prefix}: {message}")
        time.sleep(delays[attempt])


def send_table_webhook(row: dict[str, str], analysis: dict[str, Any], config: dict[str, Any]) -> int:
    webhook = table_webhook(config)
    rows = analysis_rows(row, analysis, config)
    for index, payload in enumerate(rows, start=1):
        if payload.get("分析日期"):
            payload["分析日期"] = date_millis(str(payload["分析日期"]))
        post_json(webhook["url"], payload, webhook["timeout"], "飞书表格 webhook 写入失败")
        if webhook["interval"] and index < len(rows):
            time.sleep(webhook["interval"])
    return len(rows)


def wechat_summary_text(rows: list[dict[str, str]], config: dict[str, Any]) -> str:
    lines = ["新增预警主播"]
    for index, row in enumerate(rows, start=1):
        lines.append(f"{index}.{display_value(row, '后台')}，主播昵称：{display_value(row, '主播昵称')}，抖音号：{display_value(row, '抖音号/短ID')}，运营：{display_value(row, '经纪人')}")
    table_url = str(config.get("wechat_summary", {}).get("table_url", "")).strip()
    if table_url:
        lines.append(f"表格链接：[点击打开飞书表格]({table_url})")
    return "\n".join(lines)


def send_wechat_summary(rows: list[dict[str, str]], config: dict[str, Any]) -> None:
    if not rows:
        return
    data = config.get("wechat_summary", {})
    if not data.get("enabled", False):
        return
    env = str(data.get("url_env", "WECHAT_SUMMARY_WEBHOOK_URL")).strip()
    url = str(data.get("url", "")).strip() or os.environ.get(env, "").strip()
    if not url:
        safe_print("企业微信汇总 webhook 未配置，跳过发送。")
        return
    response = requests.post(url, json={"msgtype": "markdown", "markdown": {"content": wechat_summary_text(rows, config)}}, timeout=30)
    if not response.ok:
        raise RuntimeError(f"企业微信汇总发送失败 HTTP {response.status_code}: {response.text[:500]}")
    result = response.json()
    if result.get("errcode") != 0:
        raise RuntimeError(f"企业微信汇总发送失败: {result}")
    safe_print(f"已发送企业微信新增预警主播汇总: {len(rows)} 个")
